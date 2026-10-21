import json
from datetime import date
from io import BytesIO

import pytest
from docx import Document

from contract_assistant.answer import Definition
from contract_assistant.clause_types import load_extra_clause_types
from contract_assistant.config import Settings
from contract_assistant.definitions import find_definitions
from contract_assistant.report import (
    DISCLAIMER,
    ELLIPSIS,
    MAX_QUOTE_CHARS,
    NO_DEFINITIONS,
    NO_FLAGS,
    NO_SUMMARY,
    NO_TERMS,
    Report,
    build_report,
    to_docx,
    to_markdown,
    write_report,
)
from contract_assistant.review import RULES_PATH, Flag, load_rules, review
from contract_assistant.store import read_cache
from contract_assistant.summary import Summary
from contract_assistant.term_sheet import FIELDS, LABELS, NOT_STATED, TermSheet, TermValue
from tests.conftest import ROOT
from tests.support.caches import make_cache

TAU_LOW = 0.2
TAU_HIGH = 0.8
LAW = "This Agreement is governed by the laws of Delaware."
TERM = "The Term is five years."
INSURANCE = "Each party shall carry insurance."
TEXT = f"{LAW} {TERM} {INSURANCE}"
NAME = "Supply deal"
TODAY = date(2026, 10, 21)
MODEL = "gen-test"
SUMMARY_TEXT = "A short supply deal."
TYPES = ["Governing Law", "Term", "Insurance", "Exclusivity"]
DEMO_CACHE = ROOT / "demo" / "cache" / "short_supply_agreement.json"
HEADINGS = [
    "Summary (generated)",
    "Key terms",
    "Review flags",
    "Clauses found",
    "Defined terms",
]


@pytest.fixture
def settings() -> Settings:
    return Settings()


@pytest.fixture
def cache(settings):
    return make_cache(
        TEXT,
        settings,
        {"Governing Law": [(LAW, 0.9)], "Term": [(TERM, 0.5)], "Insurance": [(INSURANCE, 0.1)]},
    )


def summary() -> Summary:
    return Summary(
        doc_id="sha256:abc",
        model=MODEL,
        tau_low=TAU_LOW,
        tau_high=TAU_HIGH,
        latency_ms=0.0,
        attempts=2,
        items=[],
        text=SUMMARY_TEXT,
    )


def sheet() -> TermSheet:
    value = TermValue(value="Delaware", clause_type="Governing Law", page=None, start=0, end=8)
    return TermSheet(
        doc_id="sha256:abc",
        model=MODEL,
        tau_low=TAU_LOW,
        tau_high=TAU_HIGH,
        prompt_version=1,
        latency_ms=0.0,
        items=[],
        terms={"governing_law": [value]},
        rejected={},
    )


def flag(unvalidated=False) -> Flag:
    return Flag(
        rule="r",
        kind="missing",
        severity="high",
        title="No cap",
        note="Exposed.",
        status="not_found",
        items=[],
        unvalidated=unvalidated,
    )


def build(cache, settings, **kwargs) -> Report:
    kwargs.setdefault("today", TODAY)
    kwargs.setdefault("tau_low", TAU_LOW)
    kwargs.setdefault("tau_high", TAU_HIGH)
    return build_report(cache, NAME, settings, TYPES, **kwargs)


def table_rows(document, index):
    return [[c.text for c in row.cells] for row in document.tables[index].rows]


def test_report_carries_title_filename_date_extractor_and_model(cache, settings):
    report = build(cache, settings, model=MODEL)
    assert report.title == NAME
    assert report.filename == "test.txt"
    assert report.made_on == "2026-10-21"
    assert report.extractor == settings.extractor_name
    assert report.model == MODEL


def test_report_size_is_in_words_for_a_text_cache(cache, settings):
    assert build(cache, settings).size == f"{len(TEXT.split())} words"


def test_report_size_is_in_pages_when_the_source_has_pages(cache, settings):
    cache.source.n_pages = 12
    assert build(cache, settings).size == "12 pages"


def test_report_made_on_defaults_to_today(cache, settings):
    made_on = build(cache, settings, today=None).made_on
    assert date.fromisoformat(made_on) >= TODAY


def test_report_thresholds_come_from_the_settings_by_default(cache, settings):
    report = build(cache, settings, tau_low=None, tau_high=None)
    assert (report.tau_low, report.tau_high) == (settings.tau_low, settings.tau_high)


def test_report_thresholds_from_the_arguments_win(cache, settings):
    report = build(cache, settings, tau_low=0.4, tau_high=0.95)
    assert (report.tau_low, report.tau_high) == (0.4, 0.95)
    assert [c.status for c in report.clauses] == ["check", "check"]


def test_report_clauses_have_one_line_per_clause_type_with_a_citation(cache, settings):
    report = build(cache, settings)
    assert [(c.clause_type, c.status, c.conf) for c in report.clauses] == [
        ("Governing Law", "found", 0.9),
        ("Term", "check", 0.5),
    ]
    assert report.clauses[0].quote == LAW
    assert report.clauses[0].page is None


def test_report_not_found_lists_the_rest_including_withheld_ones(cache, settings):
    cache.clauses["Governing Law"][0].start += 1
    report = build(cache, settings)
    assert report.not_found == ["Governing Law", "Insurance"]


def test_report_ignores_clause_types_missing_from_the_cache(cache, settings):
    report = build(cache, settings)
    assert "Exclusivity" not in report.not_found


def test_report_collapses_whitespace_in_a_quote(settings):
    quote = "Liability is capped at   fees\npaid."
    cache = make_cache(quote, settings, {"Term": [(quote, 0.9)]})
    report = build(cache, settings)
    assert report.clauses[0].quote == "Liability is capped at fees paid."


def test_report_cuts_a_long_quote_with_an_ellipsis(settings):
    quote = "word " * 100
    cache = make_cache(quote, settings, {"Term": [(quote, 0.9)]})
    shown = build(cache, settings).clauses[0].quote
    assert len(shown) == MAX_QUOTE_CHARS
    assert shown.endswith(ELLIPSIS)


def test_report_keeps_a_quote_of_exactly_the_limit(settings):
    quote = "a" * MAX_QUOTE_CHARS
    cache = make_cache(quote, settings, {"Term": [(quote, 0.9)]})
    assert build(cache, settings).clauses[0].quote == quote


def test_report_terms_are_empty_without_a_sheet(cache, settings):
    assert build(cache, settings).terms == []


def test_report_terms_have_one_line_per_field_with_a_sheet(cache, settings):
    terms = build(cache, settings, sheet=sheet()).terms
    assert [t.label for t in terms] == [LABELS[field] for field in FIELDS]
    by_label = {t.label: t for t in terms}
    assert by_label["Governing law"].value == "Delaware"
    assert by_label["Governing law"].clause_types == ["Governing Law"]
    assert by_label["Parties"].value == NOT_STATED
    assert by_label["Parties"].clause_types == []


def test_report_takes_the_summary_text_and_attempts(cache, settings):
    report = build(cache, settings, summary=summary())
    assert (report.summary, report.summary_attempts) == (SUMMARY_TEXT, 2)


def test_report_has_no_summary_when_none_is_given(cache, settings):
    report = build(cache, settings)
    assert (report.summary, report.summary_attempts) == (None, None)


def test_markdown_starts_with_the_title_line(cache, settings):
    assert to_markdown(build(cache, settings)).startswith(f"# Contract report: {NAME}\n")


def test_markdown_has_the_five_section_headings_in_order(cache, settings):
    lines = to_markdown(build(cache, settings)).splitlines()
    headings = [line[3:] for line in lines if line.startswith("## ")]
    assert headings == HEADINGS


def test_markdown_header_lines_name_file_date_models_and_thresholds(cache, settings):
    text = to_markdown(build(cache, settings, model=MODEL, tau_low=None, tau_high=None))
    assert "- File: test.txt (" in text
    assert "- Made on: 2026-10-21" in text
    assert f"language model: {MODEL}" in text
    assert f"Found from {settings.tau_high:.2f}, verify from {settings.tau_low:.2f}" in text


def test_markdown_says_none_when_no_language_model_is_given(cache, settings):
    assert "language model: none" in to_markdown(build(cache, settings))


def test_markdown_has_the_key_terms_clauses_and_defined_terms_tables(cache, settings):
    definition = Definition(term="Term", quote="Term means five years.", page=3, start=0, end=5)
    text = to_markdown(build(cache, settings, sheet=sheet(), definitions=[definition]))
    assert "| Term | Value | From clause |" in text
    assert "| Governing law | Delaware | Governing Law |" in text
    assert "| Clause | Status | Confidence | Page | Quote |" in text
    assert f"| Governing Law | found | 0.90 | n/a | {LAW} |" in text
    assert "| Term | Defined as | Page |" in text
    assert "| Term | Term means five years. | 3 |" in text


def test_markdown_shows_the_summary_and_flags(cache, settings):
    text = to_markdown(build(cache, settings, summary=summary(), flags=[flag(), flag(True)]))
    assert SUMMARY_TEXT in text
    assert "- **high** No cap: Exposed." in text
    assert "- **high** No cap (not validated on test data): Exposed." in text


def test_markdown_uses_the_empty_messages(settings):
    cache = make_cache(TEXT, settings, {})
    text = to_markdown(build_report(cache, NAME, settings, [], today=TODAY))
    for message in (NO_SUMMARY, NO_TERMS, NO_FLAGS, NO_DEFINITIONS):
        assert message in text


def test_markdown_has_a_not_found_line(cache, settings):
    assert "Not found: Insurance" in to_markdown(build(cache, settings))


def test_markdown_has_no_not_found_line_when_all_are_found(settings):
    cache = make_cache(TEXT, settings, {"Governing Law": [(LAW, 0.9)]})
    report = build_report(cache, NAME, settings, ["Governing Law"], today=TODAY)
    assert "Not found:" not in to_markdown(report)


def test_markdown_replaces_a_pipe_in_a_quote(settings):
    quote = "Pay A | B within ten days."
    cache = make_cache(quote, settings, {"Term": [(quote, 0.9)]})
    text = to_markdown(build(cache, settings))
    assert "Pay A / B within ten days." in text
    assert "A | B" not in text


def test_markdown_replaces_a_pipe_in_a_definition(cache, settings):
    definition = Definition(term="X", quote="X means a | b.", page=None, start=0, end=5)
    text = to_markdown(build(cache, settings, definitions=[definition]))
    assert "| X | X means a / b. | n/a |" in text


def test_markdown_ends_with_the_disclaimer(cache, settings):
    assert to_markdown(build(cache, settings)).rstrip("\n").endswith(DISCLAIMER)


def open_docx(report) -> Document:
    data = to_docx(report)
    assert isinstance(data, bytes)
    return Document(BytesIO(data))


def test_docx_opens_and_has_the_title_and_headings_in_order(cache, settings):
    document = open_docx(build(cache, settings))
    headings = [
        p.text for p in document.paragraphs if p.style.name.startswith(("Heading", "Title"))
    ]
    assert headings == [f"Contract report: {NAME}", *HEADINGS]


def test_docx_has_a_clauses_table_and_the_empty_message_for_missing_parts(cache, settings):
    document = open_docx(build(cache, settings))
    assert len(document.tables) == 1
    assert table_rows(document, 0)[0] == ["Clause", "Status", "Confidence", "Page", "Quote"]
    assert table_rows(document, 0)[1] == ["Governing Law", "found", "0.90", "n/a", LAW]
    texts = [p.text for p in document.paragraphs]
    assert NO_TERMS in texts
    assert NO_FLAGS in texts
    assert NO_DEFINITIONS in texts
    assert NO_SUMMARY in texts
    assert "Not found: Insurance" in texts
    assert texts[-1] == DISCLAIMER


def test_docx_has_three_tables_with_terms_and_definitions(cache, settings):
    definition = Definition(term="Term", quote="Term means five years.", page=3, start=0, end=5)
    document = open_docx(build(cache, settings, sheet=sheet(), definitions=[definition]))
    assert len(document.tables) == 3
    assert table_rows(document, 0)[0] == ["Term", "Value", "From clause"]
    assert len(table_rows(document, 0)) == len(FIELDS) + 1
    assert table_rows(document, 2)[1] == ["Term", "Term means five years.", "3"]


def test_docx_lists_flags_and_the_summary(cache, settings):
    document = open_docx(build(cache, settings, summary=summary(), flags=[flag()]))
    texts = [p.text for p in document.paragraphs]
    assert SUMMARY_TEXT in texts
    assert "high: No cap. Exposed." in texts


def test_write_report_creates_the_directory_and_both_files(cache, settings, tmp_path):
    directory = tmp_path / "out" / "reports"
    markdown_path, docx_path = write_report(build(cache, settings), directory, "deal")
    assert (markdown_path, docx_path) == (directory / "deal.md", directory / "deal.docx")
    assert markdown_path.read_text(encoding="utf-8").startswith("# Contract report")
    assert len(Document(str(docx_path)).tables) == 1


def test_a_full_report_of_the_short_demo_contract_renders_both_ways(settings):
    cache = read_cache(DEMO_CACHE, settings)
    extra = [c.name for c in load_extra_clause_types(settings.extra_clause_types_path)]
    rules = load_rules(RULES_PATH)
    flags = review(cache, rules, settings.tau_low, settings.tau_high, extra)
    definitions = find_definitions(cache.text, cache.page_starts)
    names = sorted(cache.clauses)
    report = build_report(
        cache,
        "Short supply",
        settings,
        names,
        extra,
        flags=flags,
        definitions=definitions,
        today=TODAY,
    )
    markdown = to_markdown(report)
    document = Document(BytesIO(to_docx(report)))
    assert len(flags) == 7
    assert "No cap on liability found" in markdown
    assert any("No cap on liability found" in p.text for p in document.paragraphs)
    assert json.loads(report.model_dump_json())["title"] == "Short supply"
