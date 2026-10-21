"""A report of one contract, as Markdown or DOCX: summary, key terms, flags, clauses, terms."""

from __future__ import annotations

from collections.abc import Collection, Sequence
from datetime import UTC, date, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from pydantic import BaseModel, Field

from contract_assistant.answer import STATUS_WORDS, Definition, Status, compose, verify_quotes
from contract_assistant.config import Settings
from contract_assistant.fulltext import FullTextSummary
from contract_assistant.obligations import ObligationSheet, by_party
from contract_assistant.review import Flag
from contract_assistant.store import DocumentCache
from contract_assistant.summary import Summary
from contract_assistant.term_sheet import FIELDS, LABELS, TermSheet, cell

MAX_QUOTE_CHARS = 300
ELLIPSIS = "…"
NO_PAGE = "n/a"
NO_SUMMARY = "No summary was written for this contract."
NO_TERMS = "No key terms were extracted."
NO_FLAGS = "No review flags."
NO_DEFINITIONS = "No defined terms were found."
NO_OBLIGATIONS = "No obligations were read from the clauses."
NO_FULLTEXT = "No full-text overview was written for this contract."
DISCLAIMER = (
    "Passages were located by a model and quotes were verified word for word against the "
    "contract. Text marked as generated was written by a language model from those quotes and "
    "checked against them. This is not legal advice."
)
TITLE = "Contract report"


class ClauseLine(BaseModel):
    """One clause type found in the contract, with its quote shortened for the report."""

    clause_type: str
    status: Status
    conf: float
    page: int | None
    quote: str


class TermLine(BaseModel):
    """One key term of the report."""

    label: str
    value: str
    clause_types: list[str]


class ObligationLine(BaseModel):
    """One obligation of the report."""

    party: str
    modality: str
    action: str
    clause_type: str


class Report(BaseModel):
    """Everything the report prints, assembled once so Markdown and DOCX agree."""

    title: str
    filename: str
    size: str
    made_on: str
    extractor: str
    model: str | None
    tau_low: float
    tau_high: float
    summary: str | None
    summary_attempts: int | None
    terms: list[TermLine]
    flags: list[Flag]
    clauses: list[ClauseLine]
    not_found: list[str]
    definitions: list[Definition]
    obligations: list[ObligationLine] = Field(default_factory=list)
    fulltext: str | None = None


def build_report(
    cache: DocumentCache,
    name: str,
    settings: Settings,
    clause_types: Sequence[str],
    unvalidated: Collection[str] = (),
    summary: Summary | None = None,
    sheet: TermSheet | None = None,
    flags: Sequence[Flag] = (),
    definitions: Sequence[Definition] = (),
    tau_low: float | None = None,
    tau_high: float | None = None,
    model: str | None = None,
    today: date | None = None,
    obligations: ObligationSheet | None = None,
    fulltext: FullTextSummary | None = None,
) -> Report:
    """Assemble the report of a document from its analysis and the generated parts."""
    low = settings.tau_low if tau_low is None else tau_low
    high = settings.tau_high if tau_high is None else tau_high
    present = [name for name in clause_types if name in cache.clauses]
    items, _ = verify_quotes(compose(cache, present, low, high, unvalidated), cache.text)
    clauses = [
        ClauseLine(
            clause_type=item.clause_type,
            status=item.status,
            conf=item.citation.conf,
            page=item.citation.page,
            quote=_short(item.citation.quote),
        )
        for item in items
        if item.citation is not None
    ]
    not_found = [item.clause_type for item in items if item.citation is None]
    return Report(
        title=name,
        filename=cache.source.filename,
        size=_size(cache),
        made_on=(today or datetime.now(tz=UTC).date()).isoformat(),
        extractor=settings.extractor_name,
        model=model,
        tau_low=low,
        tau_high=high,
        summary=summary.text if summary else None,
        summary_attempts=summary.attempts if summary else None,
        terms=_terms(sheet),
        flags=list(flags),
        clauses=clauses,
        not_found=not_found,
        definitions=list(definitions),
        obligations=_obligations(obligations),
        fulltext=fulltext.text if fulltext else None,
    )


def _obligations(sheet: ObligationSheet | None) -> list[ObligationLine]:
    if sheet is None:
        return []
    return [
        ObligationLine(
            party=party, modality=o.modality, action=_short(o.action), clause_type=o.clause_type
        )
        for party, obligations in by_party(sheet).items()
        for o in obligations
    ]


def _terms(sheet: TermSheet | None) -> list[TermLine]:
    if sheet is None:
        return []
    return [
        TermLine(
            label=LABELS[field],
            value=cell(sheet, field),
            clause_types=list(dict.fromkeys(v.clause_type for v in sheet.terms.get(field, []))),
        )
        for field in FIELDS
    ]


def _size(cache: DocumentCache) -> str:
    if cache.source.n_pages is not None:
        return f"{cache.source.n_pages} pages"
    return f"{len(cache.text.split()):,} words"


def _short(quote: str) -> str:
    text = " ".join(quote.split())
    if len(text) <= MAX_QUOTE_CHARS:
        return text
    return text[: MAX_QUOTE_CHARS - len(ELLIPSIS)] + ELLIPSIS


def _page(page: int | None) -> str:
    return NO_PAGE if page is None else str(page)


def to_markdown(report: Report) -> str:
    """Render the report as Markdown."""
    lines = [
        f"# {TITLE}: {report.title}",
        "",
        f"- File: {report.filename} ({report.size})",
        f"- Made on: {report.made_on}",
        f"- Extraction model: {report.extractor}; language model: {report.model or 'none'}",
        f"- Found from {report.tau_high:.2f}, verify from {report.tau_low:.2f}",
        "",
        "## Summary (generated)",
        "",
        report.summary or NO_SUMMARY,
        "",
        "## Key terms",
        "",
    ]
    if report.terms:
        lines += ["| Term | Value | From clause |", "| --- | --- | --- |"]
        lines += [
            f"| {t.label} | {t.value} | {', '.join(t.clause_types)} |" for t in report.terms
        ]
    else:
        lines.append(NO_TERMS)
    lines += ["", "## Review flags", ""]
    if report.flags:
        for flag in report.flags:
            mark = " (not validated on test data)" if flag.unvalidated else ""
            lines.append(f"- **{flag.severity}** {flag.title}{mark}: {flag.note}")
    else:
        lines.append(NO_FLAGS)
    lines += ["", "## Clauses found", "", "| Clause | Status | Confidence | Page | Quote |"]
    lines.append("| --- | --- | --- | --- | --- |")
    lines += [
        f"| {c.clause_type} | {STATUS_WORDS[c.status]} | {c.conf:.2f} | {_page(c.page)} | "
        f"{c.quote.replace('|', '/')} |"
        for c in report.clauses
    ]
    if report.not_found:
        lines += ["", f"Not found: {', '.join(report.not_found)}"]
    lines += ["", "## Obligations by party (generated, verified)", ""]
    if report.obligations:
        lines += ["| Party | | What | Clause |", "| --- | --- | --- | --- |"]
        lines += [
            f"| {o.party} | {o.modality} | {o.action.replace('|', '/')} | {o.clause_type} |"
            for o in report.obligations
        ]
    else:
        lines.append(NO_OBLIGATIONS)
    lines += ["", "## Full-text overview (generated)", "", report.fulltext or NO_FULLTEXT]
    lines += ["", "## Defined terms", ""]
    if report.definitions:
        lines += ["| Term | Defined as | Page |", "| --- | --- | --- |"]
        lines += [
            f"| {d.term} | {_short(d.quote).replace('|', '/')} | {_page(d.page)} |"
            for d in report.definitions
        ]
    else:
        lines.append(NO_DEFINITIONS)
    lines += ["", "---", "", DISCLAIMER, ""]
    return "\n".join(lines)


def to_docx(report: Report) -> bytes:
    """Render the report as a DOCX file."""
    document = Document()
    document.add_heading(f"{TITLE}: {report.title}", level=0)
    for line in (
        f"File: {report.filename} ({report.size})",
        f"Made on: {report.made_on}",
        f"Extraction model: {report.extractor}; language model: {report.model or 'none'}",
        f"Found from {report.tau_high:.2f}, verify from {report.tau_low:.2f}",
    ):
        document.add_paragraph(line)
    document.add_heading("Summary (generated)", level=1)
    document.add_paragraph(report.summary or NO_SUMMARY)
    document.add_heading("Key terms", level=1)
    if report.terms:
        _table(
            document,
            ["Term", "Value", "From clause"],
            [[t.label, t.value, ", ".join(t.clause_types)] for t in report.terms],
        )
    else:
        document.add_paragraph(NO_TERMS)
    document.add_heading("Review flags", level=1)
    if report.flags:
        for flag in report.flags:
            mark = " (not validated on test data)" if flag.unvalidated else ""
            document.add_paragraph(f"{flag.severity}: {flag.title}{mark}. {flag.note}")
    else:
        document.add_paragraph(NO_FLAGS)
    document.add_heading("Clauses found", level=1)
    _table(
        document,
        ["Clause", "Status", "Confidence", "Page", "Quote"],
        [
            [c.clause_type, STATUS_WORDS[c.status], f"{c.conf:.2f}", _page(c.page), c.quote]
            for c in report.clauses
        ],
    )
    if report.not_found:
        document.add_paragraph(f"Not found: {', '.join(report.not_found)}")
    document.add_heading("Obligations by party (generated, verified)", level=1)
    if report.obligations:
        _table(
            document,
            ["Party", "", "What", "Clause"],
            [[o.party, o.modality, o.action, o.clause_type] for o in report.obligations],
        )
    else:
        document.add_paragraph(NO_OBLIGATIONS)
    document.add_heading("Full-text overview (generated)", level=1)
    document.add_paragraph(report.fulltext or NO_FULLTEXT)
    document.add_heading("Defined terms", level=1)
    if report.definitions:
        _table(
            document,
            ["Term", "Defined as", "Page"],
            [[d.term, _short(d.quote), _page(d.page)] for d in report.definitions],
        )
    else:
        document.add_paragraph(NO_DEFINITIONS)
    document.add_paragraph(DISCLAIMER)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _table(document, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell_, header in zip(table.rows[0].cells, headers, strict=True):
        cell_.text = header
    for row in rows:
        for cell_, value in zip(table.add_row().cells, row, strict=True):
            cell_.text = value


def write_report(report: Report, directory: Path, stem: str) -> tuple[Path, Path]:
    """Write the Markdown and the DOCX of a report into a directory; return both paths."""
    directory.mkdir(parents=True, exist_ok=True)
    markdown_path = directory / f"{stem}.md"
    docx_path = directory / f"{stem}.docx"
    markdown_path.write_text(to_markdown(report), encoding="utf-8")
    docx_path.write_bytes(to_docx(report))
    return markdown_path, docx_path
